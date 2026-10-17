"""Import the universe workbook into `companies` and `archetypes`.

Priority, coverage points, tradable flag and size band are recomputed here with the workbook's own
formulas, because coverage_check.py rewrites inputs without recalculating cached formula values.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session

from localledger.models import Archetype, Company

HEADER_ROW = 4
COLUMNS = [
    "ticker",
    "name",
    "mkt_cap_musd",
    "archetype",
    "local_gov_driver",
    "public_data",
    "note",
    "analyst_coverage",
    "local_gov_intensity",
    "data_edge",
]


def coverage_points(coverage: str | None) -> float:
    return {"0 (checked)": 2.0, "Unverified": 1.0, "1 (checked)": 0.5}.get((coverage or "").strip(), 0.0)


def size_band(mkt_cap: float | None) -> str:
    cap = mkt_cap or 0.0
    if cap < 25:
        return "Nano (<$25M)"
    if cap < 100:
        return "$25-100M"
    if cap < 250:
        return "$100-250M"
    return "$250-500M"


def _num(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_universe(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Universe"]
    rows: list[dict] = []
    for idx, row in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, max_col=10, values_only=True)):
        if not row or not row[0]:
            break
        rec = dict(zip(COLUMNS, row, strict=False))
        rec["ticker"] = str(rec["ticker"]).strip().upper()
        rec["mkt_cap_musd"] = _num(rec["mkt_cap_musd"])
        rec["local_gov_intensity"] = _num(rec["local_gov_intensity"])
        rec["data_edge"] = _num(rec["data_edge"])
        rec["tradable"] = 1 if (rec["mkt_cap_musd"] or 0) >= 50 else 0
        rec["coverage_pts"] = coverage_points(rec["analyst_coverage"])
        rec["priority"] = (
            (rec["local_gov_intensity"] or 0) + (rec["data_edge"] or 0) + rec["tradable"] + rec["coverage_pts"]
        )
        rec["size_band"] = size_band(rec["mkt_cap_musd"])
        rec["sheet_rank"] = idx + 1
        rows.append(rec)
    wb.close()
    return rows


def read_archetypes(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Archetypes"]
    out = []
    for row in ws.iter_rows(min_row=4, max_col=5, values_only=True):
        if not row or not row[0]:
            continue
        out.append(
            {
                "name": str(row[0]).strip(),
                "driver": row[1],
                "datasets": row[2],
                "intensity": _num(row[3]),
                "data_edge": _num(row[4]),
            }
        )
    wb.close()
    return out


def seed_universe(session: Session, path: Path) -> tuple[int, int]:
    rows = read_universe(path)
    for rec in rows:
        existing = session.get(Company, rec["ticker"])
        if existing:
            for key, value in rec.items():
                setattr(existing, key, value)
        else:
            session.add(Company(**rec))
    archetypes = read_archetypes(path)
    for rec in archetypes:
        existing = session.get(Archetype, rec["name"])
        if existing:
            for key, value in rec.items():
                setattr(existing, key, value)
        else:
            session.add(Archetype(**rec))
    session.flush()
    return len(rows), len(archetypes)
